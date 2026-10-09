import io
import zipfile
from pathlib import Path
from .config import MAX_TEXT_CHARS, MAX_UPLOAD_BYTES

ALLOWED_EXTENSIONS = {".pdf", ".txt", ".md", ".py", ".docx"}


class FileIngester:
    @staticmethod
    def parse_bytes(filename, content):
        suffix = Path(filename).suffix.lower()
        if suffix not in ALLOWED_EXTENSIONS:
            raise ValueError("Supported formats: PDF, TXT, MD, PY, DOCX.")
        if len(content) > MAX_UPLOAD_BYTES:
            raise ValueError("Files must be smaller than 20 MB.")
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("Unlock the PDF before uploading it.")
            if len(reader.pages) > 1000:
                raise ValueError("PDFs must contain at most 1,000 pages; split the file.")
            pages = []
            total = 0
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                total += len(text)
                if total > MAX_TEXT_CHARS:
                    raise ValueError("Document text is too large; split the file.")
                pages.append({"page": i + 1, "text": text})
        elif suffix == ".docx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(entry.file_size for entry in entries) > MAX_UPLOAD_BYTES:
                    raise ValueError("DOCX expanded content is too large; split the file.")
            from docx import Document
            document = Document(io.BytesIO(content))
            parts = [p.text for p in document.paragraphs]
            parts.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
            pages = [{"page": 1, "text": "\n".join(parts)}]
        else:
            pages = [{"page": 1, "text": content.decode("utf-8-sig")}]
        if sum(len(p["text"]) for p in pages) > MAX_TEXT_CHARS:
            raise ValueError("Document text is too large; split the file.")
        if not any(p["text"].strip() for p in pages):
            raise ValueError("No readable text found. Scanned PDFs need OCR before import.")
        return pages

    @staticmethod
    async def parse_file(file):
        content = await file.read(MAX_UPLOAD_BYTES + 1)
        return "\n".join(p["text"] for p in FileIngester.parse_bytes(file.filename or "", content))

    @staticmethod
    def chunk_text(text, tokenizer=None, chunk_size=200, overlap=40):
        if chunk_size <= overlap or overlap < 0:
            raise ValueError("Overlap must be smaller than chunk size")
        offsets = None
        if tokenizer is None:
            tokens = text.split()
            decode = lambda part: " ".join(part)
        elif getattr(tokenizer, "is_fast", False):
            encoded = tokenizer(text, add_special_tokens=False, truncation=False,
                                return_offsets_mapping=True, verbose=False)
            tokens, offsets = encoded["input_ids"], encoded["offset_mapping"]
        else:
            tokens = tokenizer.encode(text, add_special_tokens=False, truncation=False)
            decode = lambda part: tokenizer.decode(part, skip_special_tokens=True)
        chunks = []
        for start in range(0, len(tokens), chunk_size - overlap):
            end = min(start + chunk_size, len(tokens))
            # Preserve source case, punctuation, whitespace and Markdown. Decoding
            # an uncased embedding tokenizer rewrites quoted evidence.
            part = (text[offsets[start][0]:offsets[end - 1][1]] if offsets is not None
                    else decode(tokens[start:end])).strip()
            if part:
                chunks.append(part)
            if start + chunk_size >= len(tokens):
                break
        return chunks
